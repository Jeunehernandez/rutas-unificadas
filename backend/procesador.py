"""
Lógica de negocio del Validador de Rutas Unificadas.

Es exactamente la misma regla del script original de escritorio,
adaptada para trabajar con bytes en memoria (subida por HTTP) en
lugar de rutas de archivo en disco y sin componentes de interfaz
gráfica (tkinter).
"""

import io
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# CONFIGURACIÓN
# ============================================================

COLUMNA_GRUPO = "Grupo"
COLUMNA_COSTO = "Costo"

NOMBRES_FECHA = [
    "Fecha",
    "Fecha Servicio",
    "Fecha del Servicio",
    "Fecha servicio",
]

NOMBRES_COLABORADOR = [
    "Titulares",
    "Colaborador",
    "Colaborador/es",
    "Empleado",
    "Empleado/a",
]

NOMBRES_RUTA = [
    "Servicio",
    "Ruta",
    "Nombre de Ruta",
]


# ============================================================
# FUNCIONES AUXILIARES
# ============================================================

def limpiar_nombre(valor):
    """Normaliza nombres de columnas."""
    if valor is None:
        return ""

    valor = str(valor).strip()
    valor = re.sub(r"\s+", " ", valor)

    return valor


def encontrar_columna(columnas, nombres_posibles):
    """
    Busca una columna utilizando coincidencia exacta
    después de normalizar los nombres.
    """

    columnas_limpias = {
        limpiar_nombre(col): col
        for col in columnas
    }

    for nombre in nombres_posibles:
        nombre_limpio = limpiar_nombre(nombre)

        if nombre_limpio in columnas_limpias:
            return columnas_limpias[nombre_limpio]

    return None


def _reiniciar_flujo(archivo):
    """Vuelve al inicio de un flujo en memoria para poder releerlo."""
    if hasattr(archivo, "seek"):
        archivo.seek(0)


def detectar_fila_encabezado(archivo):
    """
    Busca automáticamente la fila que contiene los encabezados.
    En el archivo proporcionado corresponde a la segunda fila.
    """

    _reiniciar_flujo(archivo)

    preview = pd.read_excel(
        archivo,
        sheet_name=0,
        header=None,
        nrows=15
    )

    for indice, fila in preview.iterrows():

        valores = [
            limpiar_nombre(x)
            for x in fila.tolist()
        ]

        tiene_grupo = "Grupo" in valores
        tiene_costo = "Costo" in valores

        if tiene_grupo and tiene_costo:
            return indice

    raise ValueError(
        "No se encontró una fila de encabezados que contenga "
        "'Grupo' y 'Costo'."
    )


def convertir_costo(valor):
    """
    Convierte diferentes formatos de costo a número.
    """

    if pd.isna(valor):
        return None

    if isinstance(valor, (int, float)):
        return float(valor)

    texto = str(valor).strip()

    if texto == "":
        return None

    texto = texto.replace("RD$", "")
    texto = texto.replace("$", "")
    texto = texto.replace(" ", "")

    if "," in texto and "." in texto:
        texto = texto.replace(",", "")

    elif "," in texto:
        partes = texto.split(",")

        if len(partes[-1]) == 2:
            texto = texto.replace(",", ".")
        else:
            texto = texto.replace(",", "")

    elif texto.count(".") > 1:
        texto = texto.replace(".", "")

    try:
        return float(texto)

    except ValueError:
        return None


# ============================================================
# LECTURA DEL EXCEL
# ============================================================

def cargar_excel(archivo):

    fila_encabezado = detectar_fila_encabezado(archivo)

    _reiniciar_flujo(archivo)

    df = pd.read_excel(
        archivo,
        sheet_name=0,
        header=fila_encabezado
    )

    df.columns = [
        limpiar_nombre(col)
        for col in df.columns
    ]

    df = df.dropna(axis=1, how="all")
    df = df.dropna(axis=0, how="all")

    if df.empty:
        raise ValueError(
            "El Excel no contiene registros."
        )

    return df


# ============================================================
# VALIDACIÓN
# ============================================================

def validar_columnas(df):

    columnas = df.columns.tolist()

    if COLUMNA_GRUPO not in columnas:
        raise ValueError(
            "No se encontró la columna 'Grupo'."
        )

    if COLUMNA_COSTO not in columnas:
        raise ValueError(
            "No se encontró la columna 'Costo'."
        )

    columna_fecha = encontrar_columna(columnas, NOMBRES_FECHA)

    if columna_fecha is None:
        raise ValueError(
            "No se encontró una columna de fecha."
        )

    columna_colaborador = encontrar_columna(columnas, NOMBRES_COLABORADOR)
    columna_ruta = encontrar_columna(columnas, NOMBRES_RUTA)

    return (columna_fecha, columna_colaborador, columna_ruta)


# ============================================================
# PROCESAMIENTO
# ============================================================

def procesar_datos(df, columna_fecha, columna_colaborador, columna_ruta):

    df = df.copy()

    # --------------------------------------------------------
    # FECHA
    # --------------------------------------------------------

    df["_Fecha_Analisis"] = pd.to_datetime(
        df[columna_fecha],
        dayfirst=True,
        errors="coerce"
    )

    errores_fecha = df["_Fecha_Analisis"].isna().sum()

    if errores_fecha > 0:
        raise ValueError(
            f"Se encontraron {errores_fecha} filas "
            "con fechas inválidas o vacías."
        )

    # --------------------------------------------------------
    # COSTO
    # --------------------------------------------------------

    df["_Costo_Numero"] = df[COLUMNA_COSTO].apply(convertir_costo)

    errores_costo = df["_Costo_Numero"].isna().sum()

    if errores_costo > 0:
        filas_error = df[df["_Costo_Numero"].isna()].index.tolist()

        raise ValueError(
            f"Se encontraron {errores_costo} costos "
            f"vacíos o no numéricos. "
            f"Filas: {filas_error[:10]}"
        )

    # --------------------------------------------------------
    # GRUPO
    # --------------------------------------------------------

    df["_Grupo_Limpio"] = df[COLUMNA_GRUPO].astype("string").str.strip()

    df.loc[df["_Grupo_Limpio"].eq(""), "_Grupo_Limpio"] = pd.NA

    # --------------------------------------------------------
    # CLAVE FECHA + GRUPO
    # --------------------------------------------------------

    df["_Fecha_Dia"] = df["_Fecha_Analisis"].dt.normalize()

    tiene_grupo = df["_Grupo_Limpio"].notna()

    df["Cantidad de servicios en el Grupo/Día"] = 0

    df.loc[tiene_grupo, "Cantidad de servicios en el Grupo/Día"] = (
        df.loc[tiene_grupo]
        .groupby(["_Fecha_Dia", "_Grupo_Limpio"])["_Grupo_Limpio"]
        .transform("size")
    )

    # --------------------------------------------------------
    # RUTA UNIFICADA
    # --------------------------------------------------------

    df["Es Ruta Unificada"] = False

    es_unificada = (
        tiene_grupo
        & (df["Cantidad de servicios en el Grupo/Día"] >= 2)
    )

    df.loc[es_unificada, "Es Ruta Unificada"] = True

    # --------------------------------------------------------
    # PORCENTAJE
    # --------------------------------------------------------

    df["Porcentaje Aplicado"] = 1.00

    # --------------------------------------------------------
    # IDENTIFICAR EL MAYOR COSTO
    # --------------------------------------------------------

    indices_mayor = (
        df.loc[es_unificada]
        .groupby(["_Fecha_Dia", "_Grupo_Limpio"])["_Costo_Numero"]
        .idxmax()
    )

    df.loc[es_unificada, "Porcentaje Aplicado"] = 0.50
    df.loc[indices_mayor, "Porcentaje Aplicado"] = 1.00

    # --------------------------------------------------------
    # TIPO DE COBRO
    # --------------------------------------------------------

    df["Tipo de Cobro"] = "Normal"
    df.loc[es_unificada, "Tipo de Cobro"] = "Unificada - 50%"
    df.loc[indices_mayor, "Tipo de Cobro"] = "Unificada - 100%"

    # --------------------------------------------------------
    # COSTOS
    # --------------------------------------------------------

    df["Costo Original"] = df["_Costo_Numero"]
    df["Costo Ajustado"] = df["Costo Original"] * df["Porcentaje Aplicado"]
    df["Ahorro"] = df["Costo Original"] - df["Costo Ajustado"]

    # --------------------------------------------------------
    # CAMPOS DE AUDITORÍA
    # --------------------------------------------------------

    df["Grupo"] = df["_Grupo_Limpio"]
    df["Fecha"] = df["_Fecha_Analisis"]

    if columna_colaborador:
        df["Colaborador"] = df[columna_colaborador]
    else:
        df["Colaborador"] = ""

    if columna_ruta:
        df["Ruta"] = df[columna_ruta]
    else:
        df["Ruta"] = ""

    # --------------------------------------------------------
    # DETALLE
    # --------------------------------------------------------

    columnas_a_eliminar = [
        "_Fecha_Analisis",
        "_Fecha_Dia",
        "_Costo_Numero",
        "_Grupo_Limpio",
    ]

    detalle = df.drop(columns=columnas_a_eliminar)

    return detalle


# ============================================================
# RESUMEN DIARIO
# ============================================================

def crear_resumen_diario(detalle):

    temp = detalle.copy()

    resumen = (
        temp.groupby(temp["Fecha"].dt.normalize())
        .agg(
            **{
                "Cantidad total de viajes/servicios": ("Costo Original", "size"),
                "Cantidad de rutas unificadas": ("Es Ruta Unificada", "sum"),
                "Costo original del día": ("Costo Original", "sum"),
                "Costo ajustado del día": ("Costo Ajustado", "sum"),
                "Ahorro del día": ("Ahorro", "sum"),
            }
        )
        .reset_index()
    )

    resumen = resumen.rename(columns={"Fecha": "Fecha"})

    return resumen


# ============================================================
# RESUMEN POR GRUPO
# ============================================================

def crear_resumen_grupo(detalle):

    temp = detalle[detalle["Es Ruta Unificada"] == True].copy()

    if temp.empty:
        return pd.DataFrame(
            columns=[
                "Fecha",
                "Grupo",
                "Cantidad de colaboradores/servicios",
                "Costo original",
                "Costo ajustado",
                "Ahorro",
                "Colaborador/servicio al 100 %",
                "Servicios al 50 %",
            ]
        )

    filas = []

    for (fecha, grupo), grupo_df in temp.groupby(["Fecha", "Grupo"]):

        fila_100 = grupo_df[grupo_df["Tipo de Cobro"] == "Unificada - 100%"]
        filas_50 = grupo_df[grupo_df["Tipo de Cobro"] == "Unificada - 50%"]

        colaborador_100 = "; ".join(
            fila_100["Colaborador"].fillna("").astype(str).tolist()
        )

        servicios_50 = "; ".join(
            filas_50["Colaborador"].fillna("").astype(str).tolist()
        )

        filas.append({
            "Fecha": fecha,
            "Grupo": grupo,
            "Cantidad de colaboradores/servicios": len(grupo_df),
            "Costo original": grupo_df["Costo Original"].sum(),
            "Costo ajustado": grupo_df["Costo Ajustado"].sum(),
            "Ahorro": grupo_df["Ahorro"].sum(),
            "Colaborador/servicio al 100 %": colaborador_100,
            "Servicios al 50 %": servicios_50,
        })

    return pd.DataFrame(filas)


# ============================================================
# RESUMEN MENSUAL
# ============================================================

def crear_resumen_mensual(detalle):

    total_servicios = len(detalle)

    servicios_unificados = int(detalle["Es Ruta Unificada"].sum())

    grupos_unificados = (
        detalle.loc[detalle["Es Ruta Unificada"], ["Fecha", "Grupo"]]
        .drop_duplicates()
        .shape[0]
    )

    costo_original = detalle["Costo Original"].sum()
    costo_ajustado = detalle["Costo Ajustado"].sum()
    ahorro = detalle["Ahorro"].sum()

    porcentaje_ahorro = (
        ahorro / costo_original if costo_original != 0 else 0
    )

    resumen = pd.DataFrame({
        "Concepto": [
            "Total de servicios del mes",
            "Total de servicios unificados",
            "Total de grupos/rutas unificadas",
            "Costo original total",
            "Costo ajustado total",
            "Ahorro total",
            "Porcentaje de ahorro",
        ],
        "Resultado": [
            total_servicios,
            servicios_unificados,
            grupos_unificados,
            costo_original,
            costo_ajustado,
            ahorro,
            porcentaje_ahorro,
        ]
    })

    return resumen


# ============================================================
# FORMATO EXCEL
# ============================================================

def formatear_excel(buffer_entrada):
    """
    Recibe un BytesIO con el Excel ya generado (sin formato) y
    devuelve un nuevo BytesIO con el formato aplicado.
    """

    _reiniciar_flujo(buffer_entrada)

    wb = load_workbook(buffer_entrada)

    azul = "004571"
    blanco = "FFFFFF"

    borde = Border(bottom=Side(style="thin", color="D9D9D9"))

    for ws in wb.worksheets:

        ws.freeze_panes = "A2"

        if ws.max_row >= 2:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=azul)
            cell.font = Font(bold=True, color=blanco)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = borde
                cell.alignment = Alignment(vertical="center")

        for columna in range(1, ws.max_column + 1):
            letra = get_column_letter(columna)
            max_length = 0

            for cell in ws[letra]:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))

            ws.column_dimensions[letra].width = min(max(max_length + 2, 12), 45)

    for ws in wb.worksheets:

        encabezados = {cell.value: cell.column for cell in ws[1]}

        columnas_moneda = [
            "Costo Original",
            "Costo Ajustado",
            "Ahorro",
            "Costo original",
            "Costo ajustado",
            "Costo original del día",
            "Costo ajustado del día",
            "Ahorro del día",
            "Costo original total",
            "Costo ajustado total",
            "Ahorro total",
        ]

        for nombre in columnas_moneda:
            if nombre in encabezados:
                col = encabezados[nombre]
                for fila in range(2, ws.max_row + 1):
                    ws.cell(fila, col).number_format = '"RD$" #,##0.00'

        if "Porcentaje Aplicado" in encabezados:
            col = encabezados["Porcentaje Aplicado"]
            for fila in range(2, ws.max_row + 1):
                ws.cell(fila, col).number_format = "0%"

        if "Concepto" in encabezados and "Resultado" in encabezados:
            for fila in range(2, ws.max_row + 1):
                concepto = ws.cell(fila, encabezados["Concepto"]).value

                if concepto == "Porcentaje de ahorro":
                    ws.cell(fila, encabezados["Resultado"]).number_format = "0.00%"

    for ws in wb.worksheets:

        if ws.max_row >= 2:
            nombre_tabla = "Tabla" + re.sub(r"[^A-Za-z0-9]", "", ws.title)
            ref = ws.dimensions

            tabla = Table(displayName=nombre_tabla, ref=ref)

            estilo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )

            tabla.tableStyleInfo = estilo
            ws.add_table(tabla)

    buffer_salida = io.BytesIO()
    wb.save(buffer_salida)
    buffer_salida.seek(0)

    return buffer_salida


# ============================================================
# NOMBRE DEL ARCHIVO
# ============================================================

def generar_nombre_salida(detalle):

    fecha_min = detalle["Fecha"].min()
    fecha_max = detalle["Fecha"].max()

    meses = {
        1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril",
        5: "Mayo", 6: "Junio", 7: "Julio", 8: "Agosto",
        9: "Septiembre", 10: "Octubre", 11: "Noviembre", 12: "Diciembre",
    }

    if fecha_min.month == fecha_max.month and fecha_min.year == fecha_max.year:
        nombre_mes = meses[fecha_min.month]
        return f"Transporte_{nombre_mes}_Validado.xlsx"

    return "Transporte_Validado.xlsx"


# ============================================================
# MÉTRICAS PARA LA RESPUESTA
# ============================================================

def calcular_metricas(detalle):

    costo_original = float(detalle["Costo Original"].sum())
    costo_ajustado = float(detalle["Costo Ajustado"].sum())
    ahorro = float(detalle["Ahorro"].sum())

    porcentaje_ahorro = (ahorro / costo_original) if costo_original else 0.0

    servicios_unificados = int(detalle["Es Ruta Unificada"].sum())

    grupos_unificados = (
        detalle.loc[detalle["Es Ruta Unificada"], ["Fecha", "Grupo"]]
        .drop_duplicates()
        .shape[0]
    )

    return {
        "servicios_totales": int(len(detalle)),
        "servicios_unificados": servicios_unificados,
        "grupos_unificados": grupos_unificados,
        "costo_original": costo_original,
        "costo_ajustado": costo_ajustado,
        "ahorro": ahorro,
        "porcentaje_ahorro": porcentaje_ahorro,
    }


# ============================================================
# PUNTO DE ENTRADA PRINCIPAL (usado por la API)
# ============================================================

def generar_reporte(contenido: bytes):
    """
    Recibe los bytes del Excel subido por el usuario y devuelve:
    (bytes_del_archivo_generado, nombre_sugerido, metricas)
    """

    entrada = io.BytesIO(contenido)

    df = cargar_excel(entrada)

    columna_fecha, columna_colaborador, columna_ruta = validar_columnas(df)

    detalle = procesar_datos(df, columna_fecha, columna_colaborador, columna_ruta)

    resumen_diario = crear_resumen_diario(detalle)
    resumen_grupo = crear_resumen_grupo(detalle)
    resumen_mensual = crear_resumen_mensual(detalle)

    nombre_salida = generar_nombre_salida(detalle)

    sin_formato = io.BytesIO()

    with pd.ExcelWriter(sin_formato, engine="openpyxl") as writer:
        detalle.to_excel(writer, sheet_name="Detalle", index=False)
        resumen_diario.to_excel(writer, sheet_name="Resumen Diario", index=False)
        resumen_grupo.to_excel(writer, sheet_name="Resumen por Grupo", index=False)
        resumen_mensual.to_excel(writer, sheet_name="Resumen Mensual", index=False)

    formateado = formatear_excel(sin_formato)
    metricas = calcular_metricas(detalle)

    return formateado.getvalue(), nombre_salida, metricas
